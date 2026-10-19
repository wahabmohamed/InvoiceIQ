"""Genere la presentation InvoiceIQ au format .pptx (captures d'ecran dans screenshots/)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

OUTPUT = Path(__file__).parent / "InvoiceIQ_Presentation.pptx"
SCREENSHOTS = Path(__file__).parent / "screenshots"

PRIMARY = RGBColor(0x25, 0x63, 0xEB)
DARK = RGBColor(0x1F, 0x29, 0x37)

SLIDES = [
    {
        "title": "InvoiceIQ",
        "subtitle": "Extraction intelligente et interrogation de factures par IA",
        "bullets": [],
    },
    {
        "title": "Contexte et problematique",
        "bullets": [
            "Des centaines de factures par mois (PDF scannes ou natifs)",
            "Traitement manuel : lent, couteux, source d'erreurs",
            "Besoin : automatiser l'extraction + recherche intelligente",
        ],
    },
    {
        "title": "Objectifs du projet",
        "bullets": [
            "O1 - Extraire automatiquement les donnees structurees",
            "O2 - Permettre une recherche en langage naturel",
            "O3 - Indexer plusieurs factures pour les interroger ensemble",
            "O4 - Fournir une interface web simple et intuitive",
            "O5 - Deployer via Docker en une seule commande",
        ],
    },
    {
        "title": "Fonctionnalites developpees",
        "bullets": [
            "Upload PDF/PNG/JPG + OCR (Docling)",
            "Extraction structuree (n°, date, fournisseur, HT/TVA/TTC)",
            "Indexation vectorielle (chunking + embeddings + ChromaDB)",
            "Chat : totaux exacts (SQL) + RAG pour les details",
            "Bonus : export JSON, tableau de bord",
        ],
    },
    {
        "title": "Architecture technique",
        "bullets": [
            "Frontend React <-> Backend FastAPI (API REST)",
            "Docling (OCR) + Groq (LLM) pour extraction et generation",
            "ChromaDB (vecteurs) + SQLite (metadonnees)",
            "Deploiement complet via Docker Compose",
        ],
    },
    {
        "title": "Stack technique et choix",
        "bullets": [
            "OCR : Docling (IBM, open-source)",
            "LLM : Groq (gratuit, tres faible latence)",
            "Vector DB : ChromaDB (leger, facile a dockeriser)",
            "Backend : FastAPI (async, Pydantic, doc auto)",
            "Frontend : React + TailwindCSS",
        ],
    },
    {
        "title": "Demo 1 - Factures importees et extraites",
        "image": "1_factures.png",
        "caption": "Upload -> OCR (Docling) -> extraction LLM validee : n°, fournisseur, montant TTC et devise",
    },
    {
        "title": "Demo 2 - Questions en langage naturel",
        "image": "2_chat_questions.png",
        "caption": "Total exact calcule a partir des donnees extraites + detail (chaises) retrouve par le RAG",
    },
    {
        "title": "Demo 3 - Question de suivi et extraits consultes",
        "image": "3_chat_suivi.png",
        "caption": "\"cette facture\" est reformulee en \"facture 5\" ; les extraits utilises sont affiches avec leur similarite",
    },
    {
        "title": "Demo 4 - Tableau de bord",
        "image": "4_dashboard.png",
        "caption": "Totaux par devise et par fournisseur, calcules cote serveur",
    },
    {
        "title": "Defis techniques rencontres",
        "bullets": [
            "Incompatibilite librairie groq / httpx -> version figee",
            "Modele Groq deprecie en prod -> bascule de modele",
            "Bug de volume Docker ecrasant le code source -> separation code/donnees",
            "Frontend statique sans proxy API -> reverse-proxy nginx",
            "Build Docker bloque (PyTorch CUDA ~3 Go) -> PyTorch CPU (~200 Mo)",
        ],
    },
    {
        "title": "Tests et resultats",
        "bullets": [
            "Extraction : 27/27 champs corrects (evaluation automatisee)",
            "Recherche RAG : bonne facture classee 1re pour 12/12 questions",
            "Controle comptable auto : HT + TVA = TTC, sinon 'a verifier'",
            "Mesure sur 3 factures de test : a confirmer sur des factures reelles",
        ],
    },
    {
        "title": "Conclusion et perspectives",
        "bullets": [
            "Budget total : 0 EUR (100% outils gratuits/open-source)",
            "Limites : jeu de test reduit, pas encore de vraies factures scannees",
            "Pistes : evaluation sur factures reelles, authentification, OCR multilingue",
        ],
    },
]


def add_image_slide(prs, data):
    """Slide titre + capture d'ecran centree + legende."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # "Title Only"
    slide.shapes.title.text = data["title"]
    title_par = slide.shapes.title.text_frame.paragraphs[0]
    title_par.font.color.rgb = DARK
    title_par.font.size = Pt(32)

    height = Inches(5.3)
    width = int(height * 1.6)  # captures en 16:10
    left = int((prs.slide_width - width) / 2)
    pic = slide.shapes.add_picture(str(SCREENSHOTS / data["image"]), left, Inches(1.35), width, height)
    pic.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)

    box = slide.shapes.add_textbox(Inches(0.5), Inches(6.8), prs.slide_width - Inches(1), Inches(0.5))
    par = box.text_frame.paragraphs[0]
    par.text = data["caption"]
    par.font.size = Pt(16)
    par.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    par.alignment = 2  # centre


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    # Slide de titre
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = SLIDES[0]["title"]
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(54)
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = PRIMARY
    subtitle = slide.placeholders[1]
    subtitle.text = SLIDES[0]["subtitle"]

    # Slides de contenu
    for data in SLIDES[1:]:
        if "image" in data:
            add_image_slide(prs, data)
            continue
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = data["title"]
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = DARK

        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(data["bullets"]):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = bullet
            p.font.size = Pt(22)
            p.level = 0

    prs.save(OUTPUT)
    print(f"Presentation creee : {OUTPUT}")


if __name__ == "__main__":
    build()
